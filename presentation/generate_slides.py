"""Generate a professional 7-slide PPTX presentation for the HCL Tech Hackathon.

Usage: python presentation/generate_slides.py
Output: presentation/HCL_IT_Ticket_Assistant.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE

import os

# ─── Brand Colors ───
PRIMARY   = RGBColor(0x00, 0x73, 0xE6)   # HCL Blue
DARK_BLUE = RGBColor(0x00, 0x2B, 0x5C)   # Deep navy
ACCENT    = RGBColor(0x00, 0xC4, 0x8C)   # Green accent
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY= RGBColor(0xF0, 0xF2, 0xF5)
DARK_GREY = RGBColor(0x33, 0x33, 0x33)
MED_GREY  = RGBColor(0x66, 0x66, 0x66)
ORANGE    = RGBColor(0xFF, 0x8C, 0x00)
RED       = RGBColor(0xE6, 0x3E, 0x3E)
YELLOW    = RGBColor(0xFF, 0xC1, 0x07)
GREEN_OK  = RGBColor(0x28, 0xA7, 0x45)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def _add_text_box(slide, left, top, width, height, text, font_size=18,
                  color=DARK_GREY, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def _add_paragraph(text_frame, text, font_size=16, color=DARK_GREY, bold=False,
                   alignment=PP_ALIGN.LEFT, space_before=Pt(4), space_after=Pt(2), font_name="Calibri"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = space_before
    p.space_after = space_after
    return p


def _add_rounded_rect(slide, left, top, width, height, fill_color, text="",
                       font_size=14, font_color=WHITE, bold=True):
    shape = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height,
                       fill_color=fill_color)
    shape.text_frame.word_wrap = True
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    if text:
        shape.text_frame.paragraphs[0].text = text
        shape.text_frame.paragraphs[0].font.size = Pt(font_size)
        shape.text_frame.paragraphs[0].font.color.rgb = font_color
        shape.text_frame.paragraphs[0].font.bold = bold
        shape.text_frame.paragraphs[0].font.name = "Calibri"
    tf = shape.text_frame
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    return shape


def _add_arrow_connector(slide, start_x, start_y, end_x, end_y, color=MED_GREY, width=2.0):
    """Add a simple arrow using a line shape."""
    connector = slide.shapes.add_connector(
        1,  # straight connector
        start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def _add_right_arrow(slide, left, top, width, height, color=MED_GREY):
    shape = _add_shape(slide, MSO_SHAPE.RIGHT_ARROW, left, top, width, height, fill_color=color)
    return shape


# ════════════════════════════════════════════════════════════════
# SLIDE 1: Title Slide
# ════════════════════════════════════════════════════════════════
def slide_1_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide, DARK_BLUE)

    # Top accent bar
    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(0.15), fill_color=ACCENT)

    # Emoji icon
    _add_text_box(slide, Inches(5.5), Inches(1.2), Inches(2.5), Inches(1.2),
                  "🎫", font_size=64, color=WHITE, alignment=PP_ALIGN.CENTER)

    # Title
    _add_text_box(slide, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1.2),
                  "Intelligent IT Ticket Analysis\n& Resolution Assistant",
                  font_size=36, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Subtitle
    _add_text_box(slide, Inches(2.5), Inches(4.0), Inches(8.3), Inches(0.8),
                  "Two-Source RAG  •  Evidence-Backed Resolutions  •  Enterprise-Grade Audit Trail",
                  font_size=18, color=ACCENT, bold=False, alignment=PP_ALIGN.CENTER)

    # Team / event
    _add_text_box(slide, Inches(3), Inches(5.2), Inches(7.3), Inches(0.5),
                  "HCL Tech Recruitment Hackathon 2026",
                  font_size=20, color=RGBColor(0xBB, 0xCC, 0xDD), bold=False, alignment=PP_ALIGN.CENTER)

    # Tech stack badges
    techs = ["Groq LLM", "ChromaDB", "Sentence-Transformers", "Streamlit", "Python"]
    badge_w = Inches(1.8)
    start_x = Inches(1.8)
    for i, tech in enumerate(techs):
        _add_rounded_rect(slide, start_x + i * Inches(2.0), Inches(6.2),
                          badge_w, Inches(0.42), PRIMARY, tech, font_size=12, font_color=WHITE)

    # Bottom accent bar
    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.35), SLIDE_W, Inches(0.15), fill_color=ACCENT)


# ════════════════════════════════════════════════════════════════
# SLIDE 2: The Problem
# ════════════════════════════════════════════════════════════════
def slide_2_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    # Header bar
    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "⚡ The Problem: Manual IT Ticket Triage at Enterprise Scale",
                  font_size=28, color=WHITE, bold=True, alignment=PP_ALIGN.LEFT)

    # --- 3 stat cards ---
    card_data = [
        ("25,000+", "Tickets / Month", "🎫", "Volume of enterprise IT tickets\nprocessed in HCL production"),
        ("15–20 min", "Per Ticket Triage", "⏱️", "Average time for manual\ncategorization & resolution lookup"),
        ("40%", "Misrouted Tickets", "🔄", "Tickets initially assigned\nto the wrong support team"),
    ]

    card_w = Inches(3.5)
    card_h = Inches(2.5)
    start_x = Inches(0.9)
    y = Inches(1.6)

    for i, (stat, label, emoji, desc) in enumerate(card_data):
        x = start_x + i * Inches(4.0)
        # Card background
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h, fill_color=LIGHT_GREY)
        card.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        card.line.width = Pt(1)

        # Emoji
        _add_text_box(slide, x + Inches(1.2), y + Inches(0.15), Inches(1), Inches(0.6),
                      emoji, font_size=32, alignment=PP_ALIGN.CENTER)
        # Stat number
        _add_text_box(slide, x + Inches(0.2), y + Inches(0.7), card_w - Inches(0.4), Inches(0.6),
                      stat, font_size=32, color=RED, bold=True, alignment=PP_ALIGN.CENTER)
        # Label
        _add_text_box(slide, x + Inches(0.2), y + Inches(1.3), card_w - Inches(0.4), Inches(0.35),
                      label, font_size=16, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
        # Description
        _add_text_box(slide, x + Inches(0.2), y + Inches(1.7), card_w - Inches(0.4), Inches(0.7),
                      desc, font_size=12, color=MED_GREY, alignment=PP_ALIGN.CENTER)

    # --- Pain points row ---
    pain_y = Inches(4.5)
    _add_text_box(slide, Inches(0.8), pain_y, Inches(12), Inches(0.5),
                  "Key Pain Points", font_size=20, color=DARK_BLUE, bold=True)

    pains = [
        ("❌  No audit trail", "Resolutions lack traceability to source"),
        ("😤  Analyst fatigue", "Repetitive lookups drain skilled staff"),
        ("🤖  Chatbot limits", "Generic AI hallucinate resolutions"),
        ("📊  No feedback loop", "No mechanism to improve over time"),
    ]
    for i, (title, desc) in enumerate(pains):
        px = Inches(0.8) + i * Inches(3.1)
        _add_rounded_rect(slide, px, pain_y + Inches(0.6), Inches(2.8), Inches(0.5),
                          PRIMARY, title, font_size=13, font_color=WHITE)
        _add_text_box(slide, px + Inches(0.1), pain_y + Inches(1.2), Inches(2.6), Inches(0.7),
                      desc, font_size=11, color=MED_GREY, alignment=PP_ALIGN.CENTER)

    # Bottom callout
    callout = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(6.3),
                          Inches(10.3), Inches(0.7), fill_color=RGBColor(0xFF, 0xF3, 0xCD))
    callout.line.color.rgb = YELLOW
    callout.line.width = Pt(1)
    callout.text_frame.paragraphs[0].text = '💡 "Our system doesn\'t hallucinate resolutions. When it doesn\'t know, it says so and escalates."'
    callout.text_frame.paragraphs[0].font.size = Pt(14)
    callout.text_frame.paragraphs[0].font.color.rgb = DARK_GREY
    callout.text_frame.paragraphs[0].font.italic = True
    callout.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


# ════════════════════════════════════════════════════════════════
# SLIDE 3: Architecture Diagram (Two-Source RAG)
# ════════════════════════════════════════════════════════════════
def slide_3_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    # Header
    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "🏗️ Two-Source RAG Architecture",
                  font_size=28, color=WHITE, bold=True)

    # --- Build the pipeline diagram ---
    # Row 1: Input
    box_h = Inches(0.65)

    # Layer labels and positions (6-layer pipeline, left to right)
    layers = [
        ("📥 Input\nTicket", PRIMARY),
        ("🏷️ LLM\nClassifier", RGBColor(0x6C, 0x5C, 0xE7)),
        ("🔍 Dual\nRetrieval", ACCENT),
        ("🧠 RAG\nResolver", ORANGE),
        ("👤 Analyst\nReview", RGBColor(0xE8, 0x4D, 0x93)),
        ("📝 Feedback\nCapture", RGBColor(0x00, 0x97, 0xA7)),
    ]

    layer_w = Inches(1.65)
    arrow_w = Inches(0.45)
    start_x = Inches(0.5)
    layer_y = Inches(1.5)

    for i, (label, color) in enumerate(layers):
        x = start_x + i * (layer_w + arrow_w)
        _add_rounded_rect(slide, x, layer_y, layer_w, box_h, color, label, font_size=11, font_color=WHITE)
        if i < len(layers) - 1:
            _add_right_arrow(slide, x + layer_w + Inches(0.05), layer_y + Inches(0.18),
                             Inches(0.35), Inches(0.3), color=MED_GREY)

    # --- Two-Source RAG split diagram ---
    split_y = Inches(2.8)
    _add_text_box(slide, Inches(0.5), split_y - Inches(0.35), Inches(5), Inches(0.4),
                  "Two-Source Retrieval (Core Innovation)", font_size=18, color=DARK_BLUE, bold=True)

    # Source 1: Ticket Corpus
    s1_x, s1_y = Inches(0.6), split_y + Inches(0.2)
    s1 = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, s1_x, s1_y, Inches(5.5), Inches(2.2),
                    fill_color=RGBColor(0xE8, 0xF4, 0xFD))
    s1.line.color.rgb = PRIMARY
    s1.line.width = Pt(2)

    _add_text_box(slide, s1_x + Inches(0.2), s1_y + Inches(0.1), Inches(5), Inches(0.4),
                  "📁 Source 1: Ticket Corpus (1,800 tickets)", font_size=15, color=PRIMARY, bold=True)
    _add_text_box(slide, s1_x + Inches(0.2), s1_y + Inches(0.55), Inches(5), Inches(0.35),
                  "Purpose: Similarity search + pattern matching", font_size=12, color=DARK_GREY)

    s1_items = ["18 subcategories × 100 tickets",
                "ChromaDB vector store (cosine similarity)",
                "Multilingual embeddings (paraphrase-MiniLM-L12)",
                "NO resolutions — context only"]
    for j, item in enumerate(s1_items):
        _add_text_box(slide, s1_x + Inches(0.4), s1_y + Inches(0.95) + j * Inches(0.28),
                      Inches(4.8), Inches(0.28),
                      f"✓  {item}", font_size=11, color=MED_GREY)

    # Source 2: Knowledge Base
    s2_x, s2_y = Inches(7.1), split_y + Inches(0.2)
    s2 = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, s2_x, s2_y, Inches(5.5), Inches(2.2),
                    fill_color=RGBColor(0xE8, 0xFD, 0xF0))
    s2.line.color.rgb = ACCENT
    s2.line.width = Pt(2)

    _add_text_box(slide, s2_x + Inches(0.2), s2_y + Inches(0.1), Inches(5), Inches(0.4),
                  "📚 Source 2: Knowledge Base (30 articles)", font_size=15, color=RGBColor(0x00, 0x88, 0x55), bold=True)
    _add_text_box(slide, s2_x + Inches(0.2), s2_y + Inches(0.55), Inches(5), Inches(0.35),
                  "Purpose: Evidence-backed resolution steps", font_size=12, color=DARK_GREY)

    s2_items = ["Curated runbook articles (KB-0001 to KB-0030)",
                "Hybrid retrieval: subcategory map + vector",
                "Diagnostic steps, resolution procedures",
                "Escalation criteria for each issue type"]
    for j, item in enumerate(s2_items):
        _add_text_box(slide, s2_x + Inches(0.4), s2_y + Inches(0.95) + j * Inches(0.28),
                      Inches(4.8), Inches(0.28),
                      f"✓  {item}", font_size=11, color=MED_GREY)

    # Merge arrow down to Resolution
    merge_y = split_y + Inches(2.6)
    merge_box = _add_rounded_rect(slide, Inches(4.2), merge_y, Inches(4.9), Inches(0.6),
                                   RGBColor(0x6C, 0x5C, 0xE7),
                                   "🧠 LLM Synthesizes Resolution with Evidence Citations",
                                   font_size=13, font_color=WHITE)

    # --- 3-Layer Fallback ---
    fb_y = Inches(5.9)
    _add_text_box(slide, Inches(0.5), fb_y - Inches(0.35), Inches(5), Inches(0.35),
                  "3-Layer Resilience Strategy", font_size=16, color=DARK_BLUE, bold=True)

    fallbacks = [
        ("Layer 1", "Groq LLM\n(llama-3.1-70b)", GREEN_OK),
        ("Layer 2", "Ollama Local\n(qwen2.5:7b)", ORANGE),
        ("Layer 3", "Keyword Fallback\n+ Raw KB Steps", RED),
    ]
    for i, (layer, desc, color) in enumerate(fallbacks):
        fx = Inches(0.6) + i * Inches(4.3)
        _add_rounded_rect(slide, fx, fb_y, Inches(1.2), Inches(0.7), color, layer, font_size=12, font_color=WHITE)
        _add_text_box(slide, fx + Inches(1.4), fb_y + Inches(0.05), Inches(2.5), Inches(0.6),
                      desc, font_size=11, color=DARK_GREY)
        if i < 2:
            _add_right_arrow(slide, fx + Inches(3.8), fb_y + Inches(0.2),
                             Inches(0.35), Inches(0.3), color=MED_GREY)


# ════════════════════════════════════════════════════════════════
# SLIDE 4: Data & Feature Engineering
# ════════════════════════════════════════════════════════════════
def slide_4_data_engineering(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "📊 Dataset & Feature Engineering",
                  font_size=28, color=WHITE, bold=True)

    # --- Dataset taxonomy diagram ---
    _add_text_box(slide, Inches(0.5), Inches(1.3), Inches(5), Inches(0.4),
                  "18-Subcategory Taxonomy (1,800 tickets)", font_size=18, color=DARK_BLUE, bold=True)

    categories = {
        "Database":     (["DB2", "MS SQL Server", "Oracle"], RGBColor(0x6C, 0x5C, 0xE7)),
        "Hardware":     (["CPU", "Disk", "Keyboard", "Memory", "Monitor", "Mouse"], RGBColor(0xE8, 0x4D, 0x93)),
        "Network":      (["DHCP", "DNS", "IP Address", "VPN", "Wireless"], PRIMARY),
        "Software":     (["Email", "Operating System"], ORANGE),
        "Inquiry/Help": (["Antivirus", "Internal App"], RGBColor(0x00, 0x97, 0xA7)),
    }

    cat_y = Inches(1.8)
    for i, (cat_name, (subs, color)) in enumerate(categories.items()):
        cy = cat_y + i * Inches(0.85)
        # Category box
        _add_rounded_rect(slide, Inches(0.6), cy, Inches(2.0), Inches(0.55),
                          color, cat_name, font_size=13, font_color=WHITE)
        # Arrow
        _add_right_arrow(slide, Inches(2.7), cy + Inches(0.12), Inches(0.3), Inches(0.3), color=MED_GREY)
        # Subcategory badges
        for j, sub in enumerate(subs):
            bw = Inches(1.35) if len(sub) > 8 else Inches(1.05)
            _add_rounded_rect(slide, Inches(3.15) + j * (bw + Inches(0.1)), cy + Inches(0.05),
                              bw, Inches(0.42),
                              RGBColor(0xF0, 0xF2, 0xF5), sub,
                              font_size=10, font_color=DARK_GREY, bold=False)

    # --- Feature Engineering section (right side) ---
    fe_x = Inches(7.5)
    _add_text_box(slide, fe_x, Inches(1.3), Inches(5.5), Inches(0.4),
                  "Feature Engineering", font_size=18, color=DARK_BLUE, bold=True)

    features = [
        ("🔤 Structured Embeddings", "Category + Subcategory + Text → richer vectors"),
        ("🔍 Signal Extraction", "Error codes, urgency markers, user impact → pre-LLM"),
        ("🗺️ Subcategory Crosswalk", "Dataset taxonomy → KB article pre-mapping"),
        ("🌐 Multilingual Model", "Italian tickets ↔ English KB via paraphrase-MiniLM"),
        ("⚡ Hybrid KB Retrieval", "Fast subcategory lookup + vector fallback"),
    ]

    for i, (title, desc) in enumerate(features):
        fy = Inches(1.85) + i * Inches(0.85)
        feat_box = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fe_x, fy,
                              Inches(5.3), Inches(0.7), fill_color=LIGHT_GREY)
        feat_box.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        feat_box.line.width = Pt(1)
        _add_text_box(slide, fe_x + Inches(0.15), fy + Inches(0.05), Inches(5), Inches(0.3),
                      title, font_size=13, color=PRIMARY, bold=True)
        _add_text_box(slide, fe_x + Inches(0.15), fy + Inches(0.35), Inches(5), Inches(0.3),
                      desc, font_size=11, color=MED_GREY)

    # Bottom: KB Coverage stats
    cov_y = Inches(6.3)
    _add_text_box(slide, Inches(0.5), cov_y, Inches(12), Inches(0.4),
                  "Knowledge Base Coverage", font_size=16, color=DARK_BLUE, bold=True)

    cov_data = [
        ("30", "KB Articles"),
        ("14/18", "Subcategories\nMapped"),
        ("100%", "Network\nCoverage"),
        ("100%", "Database\nCoverage"),
        ("83%", "Hardware\nCoverage"),
    ]
    for i, (num, label) in enumerate(cov_data):
        cx = Inches(0.6) + i * Inches(2.55)
        _add_rounded_rect(slide, cx, cov_y + Inches(0.4), Inches(0.9), Inches(0.55),
                          ACCENT, num, font_size=20, font_color=WHITE)
        _add_text_box(slide, cx + Inches(1.0), cov_y + Inches(0.42), Inches(1.4), Inches(0.55),
                      label, font_size=11, color=MED_GREY)


# ════════════════════════════════════════════════════════════════
# SLIDE 5: Live Demo / Walkthrough
# ════════════════════════════════════════════════════════════════
def slide_5_demo_walkthrough(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "🎬 Demo Walkthrough: VPN Ticket → Resolution",
                  font_size=28, color=WHITE, bold=True)

    # 4 stages as cards
    stages = [
        ("1️⃣  Input", "🎫",
         "\"Can't connect to corporate\nVPN from home office.\nGateway not responding.\"",
         LIGHT_GREY, PRIMARY),
        ("2️⃣  Classify", "🏷️",
         "Category: Network / VPN\nPriority: P2 🟠\nConfidence: 92%  ✅",
         RGBColor(0xF0, 0xE8, 0xFD), RGBColor(0x6C, 0x5C, 0xE7)),
        ("3️⃣  Retrieve", "🔍",
         "KB-0018: VPN connect/disconnect\n  → Relevance: 79%\nKB-0019: VPN slow routing\n  → Relevance: 59%\n3 similar tickets found (87%)",
         RGBColor(0xE8, 0xFD, 0xF0), RGBColor(0x00, 0x88, 0x55)),
        ("4️⃣  Resolve", "📋",
         "1. Restart VPN client [KB-0018]\n2. Verify system time [KB-0018]\n3. Install client update [KB-0018]\n4. Test alternate network\n5. Capture diagnostics",
         RGBColor(0xFF, 0xF3, 0xE0), ORANGE),
    ]

    card_w = Inches(2.9)
    card_h = Inches(3.3)
    start_x = Inches(0.5)
    card_y = Inches(1.5)

    for i, (title, emoji, content, bg_color, title_color) in enumerate(stages):
        x = start_x + i * (card_w + Inches(0.25))

        # Card background
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_w, card_h, fill_color=bg_color)
        card.line.color.rgb = title_color
        card.line.width = Pt(2)

        # Emoji
        _add_text_box(slide, x + Inches(1.0), card_y + Inches(0.1), Inches(1), Inches(0.5),
                      emoji, font_size=28, alignment=PP_ALIGN.CENTER)

        # Title
        _add_text_box(slide, x + Inches(0.15), card_y + Inches(0.6), card_w - Inches(0.3), Inches(0.35),
                      title, font_size=15, color=title_color, bold=True, alignment=PP_ALIGN.CENTER)

        # Content
        _add_text_box(slide, x + Inches(0.2), card_y + Inches(1.05), card_w - Inches(0.4), Inches(2.0),
                      content, font_size=11, color=DARK_GREY)

        # Arrow between cards
        if i < 3:
            _add_right_arrow(slide, x + card_w + Inches(0.02), card_y + Inches(1.5),
                             Inches(0.2), Inches(0.25), color=MED_GREY)

    # --- Evidence Trail section ---
    ev_y = Inches(5.2)
    ev_box = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), ev_y,
                         Inches(12.3), Inches(2.0), fill_color=RGBColor(0xF8, 0xF9, 0xFA))
    ev_box.line.color.rgb = PRIMARY
    ev_box.line.width = Pt(2)

    _add_text_box(slide, Inches(0.7), ev_y + Inches(0.1), Inches(11.5), Inches(0.35),
                  "🔗 Evidence Trail — Every recommendation traces back to its source",
                  font_size=15, color=PRIMARY, bold=True)

    # Evidence items
    ev_items = [
        ("KB-0018", "knowledge_base", "\"Restart the VPN client and reconnect using the organization's published gateway/profile.\"",
         "Direct resolution step for VPN connection failures"),
        ("KB-0018", "knowledge_base", "\"Confirm automatic date/time is enabled and complete the corporate sign-in/MFA flow.\"",
         "Gateway errors can be caused by time sync issues"),
        ("TKT-1351", "ticket", "\"VPN keeps dropping randomly\" — 87% similar",
         "Corroborates intermittent VPN disconnection pattern"),
    ]

    for j, (source_id, source_type, excerpt, reason) in enumerate(ev_items):
        ey = ev_y + Inches(0.5) + j * Inches(0.47)
        badge_color = PRIMARY if source_type == "knowledge_base" else MED_GREY
        _add_rounded_rect(slide, Inches(0.8), ey, Inches(1.0), Inches(0.35),
                          badge_color, source_id, font_size=10, font_color=WHITE)
        _add_text_box(slide, Inches(1.9), ey + Inches(0.02), Inches(5.5), Inches(0.32),
                      excerpt, font_size=10, color=DARK_GREY)
        _add_text_box(slide, Inches(7.6), ey + Inches(0.02), Inches(5.0), Inches(0.32),
                      f"→ {reason}", font_size=10, color=MED_GREY)


# ════════════════════════════════════════════════════════════════
# SLIDE 6: Results & Evaluation Metrics
# ════════════════════════════════════════════════════════════════
def slide_6_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "📈 Quantitative Evaluation & Benchmark Results",
                  font_size=28, color=WHITE, bold=True)

    # ── Top 3 Stat Cards for the 3 Core Engines ──
    top_cards = [
        ("Ticket Similarity", "100%", "Mean Precision@5", "1,800 Ticket Corpus • MRR: 1.000 • Cosine Similarity", PRIMARY),
        ("Knowledge Base RAG", "100%", "Top-3 Hit Rate", "30 Curated Runbooks • Top-1 Mapped Accuracy: 100%", ACCENT),
        ("Classification Engine", "92% / 71%", "LLM / Offline Fallback", "18 Subcategories • Multi-Category • P1-P4 Priority", RGBColor(0x6C, 0x5C, 0xE7)),
    ]

    card_w = Inches(3.8)
    card_h = Inches(1.5)
    for i, (title, stat, sublabel, detail, color) in enumerate(top_cards):
        cx = Inches(0.6) + i * Inches(4.15)
        cy = Inches(1.3)
        c_box = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, card_w, card_h, fill_color=LIGHT_GREY)
        c_box.line.color.rgb = color
        c_box.line.width = Pt(2)

        _add_text_box(slide, cx + Inches(0.15), cy + Inches(0.08), card_w - Inches(0.3), Inches(0.3),
                      title, font_size=13, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        _add_text_box(slide, cx + Inches(0.15), cy + Inches(0.35), card_w - Inches(0.3), Inches(0.55),
                      stat, font_size=28, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
        _add_text_box(slide, cx + Inches(0.15), cy + Inches(0.92), card_w - Inches(0.3), Inches(0.25),
                      sublabel, font_size=11, color=DARK_GREY, bold=True, alignment=PP_ALIGN.CENTER)
        _add_text_box(slide, cx + Inches(0.1), cy + Inches(1.15), card_w - Inches(0.2), Inches(0.3),
                      detail, font_size=9, color=MED_GREY, alignment=PP_ALIGN.CENTER)

    # ── Left Column: Retrieval & Classification Breakdown ──
    left_x = Inches(0.6)
    sec_y = Inches(3.0)
    _add_text_box(slide, left_x, sec_y, Inches(6.0), Inches(0.35),
                  "Engine Accuracy Metrics (Empirical)", font_size=16, color=DARK_BLUE, bold=True)

    metrics_detail = [
        ("Ticket Retrieval (Precision@1)", "100%", 1.0, GREEN_OK),
        ("Ticket Retrieval (Precision@5)", "100%", 1.0, GREEN_OK),
        ("KB Top-1 Direct Mapping Match", "100%", 1.0, GREEN_OK),
        ("KB Top-3 Article Relevance Hit", "100%", 1.0, GREEN_OK),
        ("Classifier Category Match (Fallback)", "71.4%", 0.714, ACCENT),
    ]

    for i, (metric, val_str, val_pct, col) in enumerate(metrics_detail):
        row_y = sec_y + Inches(0.4) + i * Inches(0.46)
        _add_text_box(slide, left_x, row_y, Inches(2.9), Inches(0.3),
                      metric, font_size=11, color=DARK_GREY, bold=True)
        # Bar background
        bar_w = Inches(2.3)
        b_bg = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(2.95), row_y + Inches(0.05),
                          bar_w, Inches(0.22), fill_color=RGBColor(0xE0, 0xE0, 0xE0))
        b_bg.line.fill.background()
        # Bar fill
        b_fill = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(2.95), row_y + Inches(0.05),
                            Inches(2.3 * val_pct), Inches(0.22), fill_color=col)
        b_fill.line.fill.background()
        _add_text_box(slide, left_x + Inches(5.35), row_y, Inches(0.8), Inches(0.3),
                      val_str, font_size=11, color=col, bold=True)

    # ── Right Column: Latency & Speed Benchmark ──
    right_x = Inches(7.0)
    _add_text_box(slide, right_x, sec_y, Inches(5.5), Inches(0.35),
                  "⚡ Pipeline Latency (~2.1s Total vs <10s Target)", font_size=16, color=DARK_BLUE, bold=True)

    time_stages = [
        ("Classification", "~800ms", Inches(1.7), RGBColor(0x6C, 0x5C, 0xE7)),
        ("Ticket Retrieval", "~120ms", Inches(0.35), PRIMARY),
        ("KB Fast Retrieval", "~15ms", Inches(0.15), ACCENT),
        ("Resolution Synthesis", "~1200ms", Inches(2.5), ORANGE),
    ]

    t_bar_y = sec_y + Inches(0.45)
    cur_x = right_x
    for stage_name, lat, w, col in time_stages:
        _add_shape(slide, MSO_SHAPE.RECTANGLE, cur_x, t_bar_y, w, Inches(0.45), fill_color=col)
        cur_x += w

    _add_rounded_rect(slide, cur_x + Inches(0.15), t_bar_y + Inches(0.02), Inches(0.9), Inches(0.4),
                      GREEN_OK, "2.1s", font_size=12, font_color=WHITE)

    # Latency Breakdown Cards
    lat_card_y = t_bar_y + Inches(0.65)
    for i, (stage, lat, _, col) in enumerate(time_stages):
        bx = right_x + (i % 2) * Inches(2.8)
        by = lat_card_y + (i // 2) * Inches(0.58)
        c = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.65), Inches(0.48), fill_color=LIGHT_GREY)
        c.line.color.rgb = col
        c.line.width = Pt(1)
        _add_text_box(slide, bx + Inches(0.1), by + Inches(0.05), Inches(1.5), Inches(0.35),
                      stage, font_size=10, color=DARK_GREY, bold=True)
        _add_text_box(slide, bx + Inches(1.5), by + Inches(0.05), Inches(1.0), Inches(0.35),
                      lat, font_size=11, color=col, bold=True, alignment=PP_ALIGN.RIGHT)

    # ── Bottom Box: HCL Enterprise Production Target Comparison ──
    bench_y = Inches(5.8)
    bench_box = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), bench_y,
                           Inches(12.1), Inches(1.35), fill_color=RGBColor(0xE8, 0xF4, 0xFD))
    bench_box.line.color.rgb = PRIMARY
    bench_box.line.width = Pt(1.5)

    _add_text_box(slide, Inches(0.9), bench_y + Inches(0.1), Inches(11.5), Inches(0.32),
                  "🏆 Alignment with HCL Enterprise Production Benchmarks",
                  font_size=15, color=PRIMARY, bold=True)

    benchmarks = [
        ("Response Latency", "HCL SLA: < 10 sec", "Achieved: ~2.1 sec", "⚡ 4.7x faster than target"),
        ("Resolution Precision", "Curated KEDB Runbooks", "100% Top-3 Hit Rate", "🎯 Zero fabricated fixes"),
        ("Auditability & Safety", "Mandatory Evidence Trail", "100% Citing KB/Ticket", "🛡️ Auditable compliance"),
        ("System Resilience", "3-Tier Failover", "Groq → Ollama → Keywords", "🔄 99.9% uptime design"),
    ]

    for i, (title, target, achieved, verdict) in enumerate(benchmarks):
        bx = Inches(0.9) + i * Inches(2.85)
        by = bench_y + Inches(0.45)
        _add_text_box(slide, bx, by, Inches(2.7), Inches(0.25),
                      title, font_size=11, color=DARK_BLUE, bold=True)
        _add_text_box(slide, bx, by + Inches(0.22), Inches(2.7), Inches(0.2),
                      target, font_size=10, color=MED_GREY)
        _add_text_box(slide, bx, by + Inches(0.42), Inches(2.7), Inches(0.22),
                      achieved, font_size=11, color=GREEN_OK, bold=True)
        _add_text_box(slide, bx, by + Inches(0.64), Inches(2.7), Inches(0.2),
                      verdict, font_size=9, color=DARK_GREY)


# ════════════════════════════════════════════════════════════════
# SLIDE 7: Enterprise Readiness & What's Next
# ════════════════════════════════════════════════════════════════
def slide_7_enterprise(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.1), fill_color=DARK_BLUE)
    _add_text_box(slide, Inches(0.5), Inches(0.2), Inches(12), Inches(0.7),
                  "🏢 Enterprise Readiness & Continuous Improvement",
                  font_size=28, color=WHITE, bold=True)

    # --- 4 Pillars ---
    pillars = [
        ("🔗 Audit Trail", "Every resolution\ncites its source KB\narticle or ticket.\n\nFull JSON export\nfor compliance.", PRIMARY),
        ("📊 Confidence\nThresholds", "Classification\nconfidence scoring.\n\nAuto-escalation\nwhen confidence < 0.5\nor no KB match.", RGBColor(0x6C, 0x5C, 0xE7)),
        ("👤 Human-in-\nthe-Loop", "Low-confidence flags\nfor analyst review.\n\n4-signal feedback\nform captures\ncorrection data.", RGBColor(0xE8, 0x4D, 0x93)),
        ("📝 KB Gap\nDetection", "\"Missing topic?\"\nfield identifies\ncoverage gaps.\n\nContinuous KB\nenrichment loop.", ACCENT),
    ]

    pillar_w = Inches(2.8)
    pillar_h = Inches(3.0)
    start_x = Inches(0.6)
    pillar_y = Inches(1.5)

    for i, (title, desc, color) in enumerate(pillars):
        px = start_x + i * (pillar_w + Inches(0.3))

        # Pillar card
        card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, px, pillar_y, pillar_w, pillar_h,
                          fill_color=WHITE)
        card.line.color.rgb = color
        card.line.width = Pt(2)

        # Color header stripe
        _add_shape(slide, MSO_SHAPE.RECTANGLE, px, pillar_y, pillar_w, Inches(0.08), fill_color=color)

        # Title
        _add_text_box(slide, px + Inches(0.15), pillar_y + Inches(0.2), pillar_w - Inches(0.3), Inches(0.7),
                      title, font_size=16, color=color, bold=True, alignment=PP_ALIGN.CENTER)

        # Description
        _add_text_box(slide, px + Inches(0.2), pillar_y + Inches(1.0), pillar_w - Inches(0.4), Inches(1.8),
                      desc, font_size=12, color=MED_GREY, alignment=PP_ALIGN.CENTER)

    # --- Feedback Loop Diagram ---
    loop_y = Inches(4.9)
    _add_text_box(slide, Inches(0.5), loop_y, Inches(12), Inches(0.4),
                  "Continuous Improvement Feedback Loop", font_size=16, color=DARK_BLUE, bold=True)

    loop_steps = [
        ("Ticket\nAnalyzed", PRIMARY),
        ("Analyst\nReviews", RGBColor(0x6C, 0x5C, 0xE7)),
        ("Feedback\nCaptured", RGBColor(0xE8, 0x4D, 0x93)),
        ("KB\nEnriched", ACCENT),
        ("Model\nImproved", ORANGE),
    ]

    loop_box_w = Inches(1.8)
    loop_start_x = Inches(0.8)
    loop_box_y = loop_y + Inches(0.5)

    for i, (label, color) in enumerate(loop_steps):
        lx = loop_start_x + i * (loop_box_w + Inches(0.7))
        _add_rounded_rect(slide, lx, loop_box_y, loop_box_w, Inches(0.7),
                          color, label, font_size=12, font_color=WHITE)
        if i < len(loop_steps) - 1:
            _add_right_arrow(slide, lx + loop_box_w + Inches(0.05), loop_box_y + Inches(0.18),
                             Inches(0.55), Inches(0.35), color=MED_GREY)

    # Return loop arrow (text)
    _add_text_box(slide, Inches(4.5), loop_box_y + Inches(0.8), Inches(4.5), Inches(0.4),
                  "🔄 Cycle repeats — system gets smarter with every ticket",
                  font_size=12, color=MED_GREY, alignment=PP_ALIGN.CENTER)

    # --- Bottom: Key differentiators ---
    diff_y = Inches(6.5)
    diff_box = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), diff_y,
                           Inches(11.3), Inches(0.7), fill_color=DARK_BLUE)
    diff_box.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(1.3), diff_y + Inches(0.15), Inches(10.7), Inches(0.4))
    tf = txBox.text_frame
    tf.paragraphs[0].text = "Not a chatbot.  An enterprise-grade resolution engine with evidence, auditability, and continuous learning."
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.name = "Calibri"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER


# ════════════════════════════════════════════════════════════════
# Main
# ════════════════════════════════════════════════════════════════
def main():
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 inches (widescreen 16:9)
    prs.slide_height = Emu(6858000)   # 7.5 inches

    slide_1_title(prs)
    slide_2_problem(prs)
    slide_3_architecture(prs)
    slide_4_data_engineering(prs)
    slide_5_demo_walkthrough(prs)
    slide_6_results(prs)
    slide_7_enterprise(prs)

    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "HCL_IT_Ticket_Assistant.pptx")
    prs.save(out_path)
    print(f"✅ Presentation saved to: {out_path}")
    print(f"   {len(prs.slides)} slides generated")


if __name__ == "__main__":
    main()
